"""Read creator CSV/XLSX exports as data; never execute formulas or macros."""
import csv
import io
from pathlib import Path
from .base import ProviderError


def read_export(path, *, sheet=None):
    path = Path(path)
    if path.stat().st_size > 30 * 1024 * 1024:
        raise ProviderError("export_too_large", "导出文件超过 30 MB，请缩小导出范围")
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        text = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ProviderError("export_format", "无法识别 CSV 编码")
        table = list(csv.reader(io.StringIO(text)))
    elif path.suffix.lower() == ".xlsx":
        import zipfile
        with zipfile.ZipFile(path) as archive:
            if sum(x.file_size for x in archive.infolist()) > 150 * 1024 * 1024:
                raise ProviderError("export_too_large", "工作簿展开后过大")
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ProviderError("export_format", "多工作表导出需要明确指定 sheet")
            tab = book[sheet] if sheet else book.active
            table = list(tab.iter_rows(values_only=True))
        finally:
            book.close()
    else:
        raise ProviderError("export_format", "目前支持 CSV 和 XLSX 导出")
    if not table:
        raise ProviderError("export_format", "导出文件缺少表头")
    headers = [str(x or "").strip() for x in table[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ProviderError("export_format", "表头为空或重复，无法可靠映射")
    rows = []
    for values in table[1:]:
        if not any(x not in (None, "") for x in values):
            continue
        if len(values) != len(headers):
            raise ProviderError("export_format", "行列数量与表头不一致")
        rows.append(dict(zip(headers, values)))
    return rows
